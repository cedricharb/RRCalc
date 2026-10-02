import argparse
import json
import re
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SHEETS = [
    "Kanto Leaders",
    "Kanto Rematch",
    "Johto Leaders",
    "Rivals",
    "Team Rocket",
    "Mini Bosses",
    "Optional Bosses",
    "Indigo League",
    "Postgame",
]

SPECIES_COLS = [5, 10, 15, 20, 25, 30]

KANTO_GYM_LOCATIONS = {
    "BROCK": "PEWTER CITY",
    "MISTY": "CERULEAN CITY",
    "LT. SURGE": "VERMILLION CITY",
    "ERIKA": "CELADON CITY",
    "SABRINA": "SAFFRON CITY",
    "KOGA": "FUCHSIA CITY",
    "BLAINE": "CINNABAR ISLAND",
    "CLAIR": "VIRIDIAN CITY",
}

LOCATION_ALIASES = {
    "VIRID. FOREST": "VIRIDIAN FOREST",
    "NUGG. BRIDGE": "NUGGET BRIDGE",
    "VERMILLION GYM": "VERMILLION CITY GYM",
    "CELADON GYM": "CELADON CITY GYM",
    "FUCHSIA GYM": "FUCHSIA CITY GYM",
    "PKMN TOWER": "POKEMON TOWER",
    "PKM MANSION": "POKEMON MANSION",
    "MANSION": "POKEMON MANSION",
    "MANSION F4": "POKEMON MANSION 4F",
    "ROCKET HIDE.": "ROCKET HIDEOUT",
    "CERULEA. CAVE": "CERULEAN CAVE",
    "SEAFOAM ISL.": "SEAFOAM ISLANDS",
    "CINNABAR ISL.": "CINNABAR ISLAND",
    "CINNABAR GYM": "CINNABAR ISLAND GYM",
    "CINNABAR LAB": "CINNABAR ISLAND LAB",
    "INDIGO PLAT.": "INDIGO PLATEAU",
    "TREASURE BEA.": "TREASURE BEACH",
    "PEWTER MUSEUM": "PEWTER CITY",
    "SAFFRON DOJO": "SAFFRON CITY",
    "ROUTE 22 #1": "ROUTE 22",
    "ROUTE 22 #2": "ROUTE 22",
    "CELADON HOTEL": "CELADON CITY HOTEL",
    "JOYFUL GAME CORNER": "JOYFUL GAME CORNER",
    "POWER PLANT": "POWER PLANT",
    "ROCK TUNNEL": "ROCK TUNNEL",
    "LAVENDER TOWN": "LAVENDER TOWN",
    "MT. MOON": "MT. MOON",
    "MT. EMBER": "MT. EMBER",
    "S.S. ANNE": "S.S. ANNE",
    "SILPH CO.": "SILPH CO.",
    "SEAFOAM ISLANDS SECRET ROOM": "SEAFOAM ISLANDS SECRET ROOM",
}

ROLE_OMIT = {"LEADER", "GYM LEADER", "ELITE FOUR", "CHAMPION"}
PAIR_DROP_ROLES = {
    "ACE TRAINER",
    "BEAUTY",
    "BUG CATCHER",
    "COOL COUPLE",
    "CRUSH KIN",
    "DUMASS",
    "GENTLEMAN",
    "LASS",
    "SAILOR",
    "TWINS",
}

ORDER_NAME_ALIASES = {
    "ACE ": "ACE TRAINER ",
    "CATCHER ": "BUG CATCHER ",
    "MANIAC ": "POKE MANIAC ",
    "BIRD KE. ": "BIRD KEEPER ",
    "TAMER ": "DRAGON TAMER ",
    "ACE NELLE": "ACE TRAINER NELLE",
    "ACE MARY": "ACE TRAINER MARY",
    "ACE COLE": "ACE TRAINER COLE",
    "ACE HALEY": "ACE TRAINER HALEY",
    "ACE WILTON": "ACE TRAINER WILTON",
    "ACE NAOMI": "ACE TRAINER NAOMI",
    "ACE ROLANDO": "ACE TRAINER ROLANDO",
    "ACE GEORGE": "ACE TRAINER GEORGE",
    "ACE CAROLINE": "ACE TRAINER CAROLINE",
    "ACE COLBY": "ACE TRAINER COLBY",
    "ACE ALEXA": "ACE TRAINER ALEXA",
    "ACE DEREK": "ACE TRAINER DEREK",
    "ACE LUCY": "ACE TRAINER LUCY",
    "ACE ZAC": "ACE TRAINER ZAC",
}

BLOCK_LABEL_ALIASES = {
    "LASS RELI - NUGGET BRIDGE": "LASS RELLI - NUGGET BRIDGE",
}

SPECIES_ALIASES = {
    "aegislash": "aegislashblade",
    "arcanineh": "arcaninehisui",
    "articunog": "articunogalar",
    "avaluggh": "avalugghisui",
    "braviaryh": "braviaryhisui",
    "calyrexi": "calyrexice",
    "calyrexs": "calyrexshadow",
    "centiskorchmegas": "centiskorchseviimega",
    "clawitzers": "clawitzersevii",
    "cramorantgorg": "cramorantgorging",
    "darmanitangz": "darmanitangalarzen",
    "darmanitanz": "darmanitanzen",
    "decidueyeh": "decidueyehisui",
    "deoxysa": "deoxysattack",
    "dodrios": "dodriosevii",
    "dugtrioa": "dugtrioalola",
    "electrodeh": "electrodehisui",
    "enamorusi": "enamorus",
    "enamorust": "enamorustherian",
    "eternatusmax": "eternatus",
    "farfetchdg": "farfetchdgalar",
    "giratinao": "giratinaorigin",
    "goodrah": "goodrahisui",
    "groudonp": "groudonprimal",
    "hoopau": "hoopaunbound",
    "indeedeem": "indeedee",
    "kyogrep": "kyogreprimal",
    "kyuremb": "kyuremblack",
    "kyuremw": "kyuremwhite",
    "landorusi": "landorus",
    "landorust": "landorustherian",
    "lilliganth": "lilliganthisui",
    "lycanrocd": "lycanrocdusk",
    "magearnao": "magearnaoriginal",
    "mantines": "mantinesevii",
    "marowaka": "marowakalola",
    "milotics": "miloticsevii",
    "moltresg": "moltresgalar",
    "mrmimeg": "mrmimegalar",
    "muka": "mukalola",
    "necrozmadm": "necrozmaduskmane",
    "ninetalesa": "ninetalesalola",
    "ogerponc": "ogerponcornerstone",
    "ogerponh": "ogerponhearthflame",
    "ogerponw": "ogerponwellspring",
    "palkiao": "palkiaorigin",
    "persiana": "persianalola",
    "qwilfishh": "qwilfishhisui",
    "raichua": "raichualola",
    "rotomf": "rotomfrost",
    "rotomh": "rotomheat",
    "rotomw": "rotomwash",
    "samurotth": "samurotthisui",
    "sandslasha": "sandslashalola",
    "shaymins": "shayminsky",
    "slowbrog": "slowbrogalar",
    "sneaselh": "sneaselhisui",
    "squawkabillyg": "squawkabilly",
    "thundurust": "thundurustherian",
    "tornadusi": "tornadus",
    "ursalunabm": "ursalunabloodmoon",
    "ursarings": "ursaringsevii",
    "urshifur": "urshifurapidstrike",
    "urshifus": "urshifu",
    "voltorbh": "voltorbhisui",
    "weezingg": "weezinggalar",
    "wishiwashisch": "wishiwashischool",
    "wishiwashissch": "wishiwashiseviischool",
    "wormadamsa": "wormadamsandy",
    "zacianc": "zaciancrowned",
    "zamazentac": "zamazentacrowned",
    "zapdosg": "zapdosgalar",
    "zebstrikas": "zebstrikasevii",
    "zoroarkh": "zoroarkhisui",
    "zygardec": "zygardecomplete",
}

MOVE_ALIASES = {
    "behembash": "behemothbash",
    "behemblade": "behemothblade",
    "behemobash": "behemothbash",
    "ceaseedge": "ceaselessedge",
    "clangscales": "clangingscales",
    "dazzgleam": "dazzlinggleam",
    "diamstorm": "diamondstorm",
    "doubleibash": "doubleironbash",
    "dragonhamm": "dragonhammer",
    "drainkiss": "drainingkiss",
    "dynacannon": "dynamaxcannon",
    "eathquake": "earthquake",
    "expandforce": "expandingforce",
    "firstimpress": "firstimpression",
    "gigahammer": "gigatonhammer",
    "highhorsep": "highhorsepower",
    "highjkick": "highjumpkick",
    "hyperspfury": "hyperspacefury",
    "mistyexplosi": "mistyexplosion",
    "parabcharge": "paraboliccharge",
    "populabomb": "populationbomb",
    "powuppunch": "poweruppunch",
    "precipblades": "precipiceblades",
    "scorchsands": "scorchingsands",
    "steamerupt": "steameruption",
    "stomtantrum": "stompingtantrum",
    "surgstrikes": "surgingstrikes",
    "thousarrows": "thousandarrows",
    "thundkick": "thunderouskick",
    "watershurik": "watershuriken",
}

ITEM_ALIASES = {
    "abomasnite": "abomasite",
    "applite": "flapplite",
    "cornermask": "cornerstonemask",
    "eternamaxorb": "eternamaxorb",
    "hearthmask": "hearthflamemask",
    "weaknesspol": "weaknesspolicy",
    "wellmask": "wellspringmask",
}

HIDDEN_POWER_IVS = {
    "BUG": {"at": 30, "df": 30, "sd": 30},
    "DRAGON": {"at": 30},
    "ELECTRIC": {"sa": 30},
    "FIGHTING": {"df": 30, "sa": 30, "sd": 30, "sp": 30},
    "FIRE": {"at": 30, "sa": 30, "sp": 30},
    "FLYING": {"hp": 30, "at": 30, "df": 30, "sa": 30, "sd": 30},
    "GHOST": {"df": 30, "sd": 30},
    "GRASS": {"at": 30, "sa": 30},
    "GROUND": {"sa": 30, "sd": 30},
    "ICE": {"at": 30, "df": 30},
    "POISON": {"df": 30, "sa": 30, "sd": 30},
    "PSYCHIC": {"at": 30, "sp": 30},
    "ROCK": {"df": 30, "sd": 30, "sp": 30},
    "STEEL": {"sd": 30},
    "WATER": {"at": 30, "df": 30, "sa": 30},
}


def flat(value):
    return " ".join(str(value).replace("\n", " ").split()).strip().upper()


def title_lines(value):
    return [flat(line) for line in str(value).split("\n") if flat(line)]


def to_id(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def load_canonical_maps(repo_root):
    node_script = r"""
const fs = require('fs');
const vm = require('vm');
const {SPECIES, MOVES, ITEMS, ABILITIES, toID} = require('./calc/dist');

function addObjectEntries(target, source) {
  for (const genData of source) {
    if (!genData || typeof genData !== 'object') continue;
    for (const key of Object.keys(genData)) {
      const value = genData[key];
      const name = value && value.name ? value.name : key;
      target[toID(name)] = name;
    }
  }
}

function addStringEntries(target, source) {
  for (const genData of source) {
    if (!Array.isArray(genData)) continue;
    for (const name of genData) {
      target[toID(name)] = name;
    }
  }
}

const maps = {species: {}, moves: {}, items: {}, abilities: {}};
addObjectEntries(maps.species, SPECIES);
addObjectEntries(maps.moves, MOVES);
addStringEntries(maps.items, ITEMS);
addStringEntries(maps.abilities, ABILITIES);

const source = fs.readFileSync('./src/js/data/sets/hardcore.js', 'utf8');
const sandbox = {};
vm.createContext(sandbox);
vm.runInContext(source + ';this.__setdex = SETDEX_SV;', sandbox);
const setdex = sandbox.__setdex;
for (const species of Object.keys(setdex)) {
  maps.species[toID(species)] = species;
  for (const set of Object.values(setdex[species])) {
    if (set.ability) maps.abilities[toID(set.ability)] = set.ability;
    if (set.item) maps.items[toID(set.item)] = set.item;
    for (const move of set.moves || []) {
      maps.moves[toID(move)] = move;
    }
  }
}

console.log(JSON.stringify(maps));
"""
    result = subprocess.run(
        ["node", "-e", node_script],
        cwd=repo_root,
        capture_output=True,
        check=True,
    )
    return json.loads(result.stdout.decode("utf-8"))


def canonicalize_order_name(text):
    name = flat(text)
    if name in ORDER_NAME_ALIASES:
        return ORDER_NAME_ALIASES[name]
    for prefix, replacement in ORDER_NAME_ALIASES.items():
        if prefix.endswith(" ") and name.startswith(prefix):
            return replacement + name[len(prefix):]
    return name


def normalize_location(text):
    return LOCATION_ALIASES.get(flat(text), flat(text))


def parse_cap_value(value):
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value)
    match = re.search(r"([0-9]+)\s*/\s*([0-9]+)", text)
    if match:
        return int(match.group(2))
    match = re.search(r"([0-9]+)\s*\+", text)
    if match:
        return int(match.group(1))
    match = re.search(r"([0-9]+)", text)
    return int(match.group(1)) if match else None


def cell_fill_color(cell):
    fill = cell.fill
    if not fill or not fill.fill_type:
        return None
    color = fill.fgColor
    if color.type == "rgb":
        return color.rgb
    if color.type == "indexed":
        return f"indexed:{color.indexed}"
    if color.type == "theme":
        return f"theme:{color.theme}:tint:{color.tint}"
    return color.type


def assign_trainer_sides(members):
    colors = []
    for member in members:
        color = member.get("fillColor")
        if color and color not in colors:
            colors.append(color)
    if len(colors) < 2:
        return
    for member in members:
        member["trainerSide"] = "left" if member.get("fillColor") == colors[0] else "right"


def iter_order_entries(workbook):
    sheet = workbook["Trainer Order"]
    entries = []
    rival_route_22_count = 0
    for row in range(1, sheet.max_row):
        trainer = sheet.cell(row, 4).value
        location = sheet.cell(row + 1, 4).value
        if not isinstance(trainer, str) or not isinstance(location, str):
            continue
        trainer_label = flat(trainer)
        if trainer_label.startswith("("):
            continue
        name = canonicalize_order_name(trainer_label)
        loc = normalize_location(location)
        label = f"{name} - {loc}"
        if label == "RIVAL - ROUTE 22":
            rival_route_22_count += 1
            if rival_route_22_count == 2:
                label = "RIVAL 4 - ROUTE 22"
        entries.append({"label": label, "cap": parse_cap_value(sheet.cell(row, 6).value)})
    return entries


def extract_note_tags(sheet, start_row):
    notes = []
    for row in range(max(1, start_row - 6), start_row):
        for col in (3, 5):
            value = sheet.cell(row, col).value
            if isinstance(value, str):
                text = flat(value)
                if "(!)" in text or "BACK TO BACK" in text:
                    notes.append(text)
    return notes


def note_variant(notes):
    combined = " | ".join(notes)
    if "IF RIVAL HAS SQUIRTLE" in combined:
        return "SQUIRTLE"
    if "IF RIVAL HAS CHARMANDER" in combined:
        return "CHARMANDER"
    if "IF RIVAL HAS BULBASAUR" in combined:
        return "BULBASAUR"
    if "PRE LT. SURGE" in combined:
        return "PRE-SURGE"
    if "POST LT. SURGE" in combined:
        return "POST-SURGE"
    if "RAIN TEAM" in combined:
        return "RAIN"
    if "SNOW TEAM" in combined:
        return "SNOW"
    if "TEAM ONE" in combined:
        return "TEAM 1"
    if "TEAM TWO" in combined:
        return "TEAM 2"
    if "TEAM THREE" in combined:
        return "TEAM 3"
    if "BACK TO BACK" in combined:
        return "BACK TO BACK"
    return None


def role_name(parts):
    if not parts:
        return ""
    if "&" in parts:
        names = [part for part in parts if part != "&" and part not in ROLE_OMIT and part not in PAIR_DROP_ROLES]
        return " & ".join(names)
    if parts[0] in ROLE_OMIT:
        return " ".join(parts[1:])
    return " ".join(parts)


def build_base_label(sheet_name, title, notes):
    parts = title_lines(title)
    if not parts or parts[0] == "PARTNER":
        return None

    if sheet_name == "Kanto Leaders":
        name = parts[-1]
        return f"{name} - {KANTO_GYM_LOCATIONS[name]}"

    if sheet_name == "Kanto Rematch":
        name = parts[-1]
        return f"{name} (REMATCH) - {KANTO_GYM_LOCATIONS[name]}"

    if sheet_name == "Johto Leaders":
        return f"{parts[-1]} - {normalize_location(parts[0])}"

    if sheet_name == "Rivals":
        if parts[-1] == "RIVAL":
            if parts[0] == "ROUTE 22 #2":
                return "RIVAL 4 - ROUTE 22"
            if parts[0] == "CHAMPION":
                return "RIVAL - CHAMPION"
            return f"RIVAL - {normalize_location(parts[0])}"
        if parts[-1] == "BRENDAN" and parts[0] == "FUCHSIA CITY":
            return "BRENDAN - SAFARI ZONE"
        return f"{parts[-1]} - {normalize_location(parts[0])}"

    if sheet_name == "Team Rocket":
        if parts[0] == "SILPH CO." and parts[1:] == ["ARIANA", "&", "ARCHER"]:
            return "ARIANA & ARCHER - SILPH CO."
        return f"{role_name(parts[1:])} - {normalize_location(parts[0])}"

    if sheet_name in {"Mini Bosses", "Optional Bosses"}:
        label = f"{role_name(parts[1:])} - {normalize_location(parts[0])}"
        if label == "CREATOR SOUPERCELL - INDIGO PLATEAU":
            label = "DUMASS CREATOR - INDIGO PLATEAU"
        return BLOCK_LABEL_ALIASES.get(label, label)

    if sheet_name == "Indigo League":
        if parts[-1] == "RIVAL":
            return "RIVAL - CHAMPION"
        return f"{parts[-1]} - {'ELITE FOUR' if parts[0] == 'ELITE FOUR' else 'CHAMPION'}"

    if sheet_name == "Postgame":
        if parts[0] == "GYM LEADER":
            name = parts[-1]
            mapped = {"CLAIR": "VIRIDIAN CITY", "MISTY": "CERULEAN CITY"}[name]
            return f"{name} - {mapped} [POSTGAME]"
        if parts[0] == "OAK'S LAB":
            return "PROFESSOR OAK - OAK'S LAB"
        if parts[0] == "PKMN TOWER" and len(parts) >= 4 and parts[1].startswith("GAUNTLET"):
            return f"{role_name(parts[2:])} - POKEMON TOWER [{parts[1]}]"
        if parts[0] == "SEAFOAM ISL." and len(parts) >= 4 and parts[1] == "SECRET ROOM":
            return "PRYCE - SEAFOAM ISLANDS SECRET ROOM"
        if parts[0] == "JOYFUL" and len(parts) >= 4 and parts[1] == "GAME CORNER":
            return f"{role_name(parts[2:])} - JOYFUL GAME CORNER"
        if parts[0] == "BERRY FOREST" and len(parts) >= 4 and parts[1].startswith("GAUNTLET"):
            return f"{role_name(parts[2:])} - BERRY FOREST [{parts[1]}]"
        return f"{role_name(parts[1:])} - {normalize_location(parts[0])}"

    return None


def is_block_row(sheet, row):
    title = sheet.cell(row, 3).value
    species = sheet.cell(row, 5).value
    if not isinstance(title, str) or not isinstance(species, str):
        return False
    title_label = flat(title)
    species_label = flat(species)
    if (
        title_label.startswith("IF YOU'RE")
        or title_label.startswith("LEVEL ")
        or title_label.startswith("(!!!)")
        or title_label.startswith("(!)")
        or title_label.startswith("(OPTIONAL)")
    ):
        return False
    return species_label not in {"BASE STATS", "HP", "ATK", "DEF", "SPA", "SPD", "SPE", "SPEED STAT:"}


def get_block_cap(sheet_name, base_label, order_caps):
    if sheet_name == "Indigo League":
        return 85
    if sheet_name == "Postgame":
        return 100
    return order_caps.get(base_label)


def parse_level(sheet, sheet_name, base_label, block_row, species_col, order_caps):
    raw = sheet.cell(block_row + 1, species_col).value
    if isinstance(raw, (int, float)):
        return int(raw)

    if isinstance(raw, str):
        raw_label = flat(raw)
        if raw_label.startswith("HIGHEST LV") or raw_label == "PLAYER MAX LEVEL":
            base_cap = get_block_cap(sheet_name, base_label, order_caps)
            if base_cap is not None:
                match = re.search(r"-\s*([0-9]+)", raw_label)
                return base_cap - int(match.group(1)) if match else base_cap

    formula = sheet.cell(block_row + 19, species_col + 3).value
    if isinstance(formula, str):
        match = re.search(r"\*([0-9]+(?:\.[0-9]+)?)\)/100\+5", formula)
        if match:
            return int(float(match.group(1)))

    if isinstance(raw, str):
        match = re.search(r"([0-9]+)", raw)
        if match:
            return int(match.group(1))

    raise ValueError(f"Unable to parse level at row {block_row}, col {species_col}")


def resolve_name(raw, canonical_map, alias_map=None):
    text = " ".join(str(raw).split()).strip()
    key = to_id(text)
    if key in canonical_map:
        return canonical_map[key]
    if alias_map and key in alias_map:
        alias_key = alias_map[key]
        if alias_key in canonical_map:
            return canonical_map[alias_key]
        raise KeyError(f"Alias target {alias_key!r} missing for {text!r}")
    raise KeyError(text)


def normalize_species(raw, canonical_maps):
    return resolve_name(raw, canonical_maps["species"], SPECIES_ALIASES)


def normalize_move(raw, canonical_maps):
    text = " ".join(str(raw).split()).strip()
    upper_text = text.upper()
    if upper_text.startswith("HP "):
        hp_type = upper_text.split(" ", 1)[1]
        hidden_power = f"Hidden Power {hp_type.title()}"
        return resolve_name(hidden_power, canonical_maps["moves"])
    return resolve_name(text, canonical_maps["moves"], MOVE_ALIASES)


def normalize_item(raw, canonical_maps):
    text = " ".join(str(raw).split()).strip()
    if not text or flat(text) == "NO ITEM":
        return None
    return resolve_name(text, canonical_maps["items"], ITEM_ALIASES)


def normalize_ability(raw, species_name, canonical_maps):
    lines = [line.strip() for line in str(raw).splitlines() if line.strip()]
    species_id = to_id(species_name)
    if species_id in {"calyrexice", "calyrexshadow"}:
        return resolve_name(
            "As One (Glastrier)" if species_id == "calyrexice" else "As One (Spectrier)",
            canonical_maps["abilities"],
        )

    choice = lines[0] if lines else ""
    if len(lines) > 1:
        for line in reversed(lines):
            line_upper = flat(line)
            if (
                ("MEGA" in line_upper and "mega" in species_id)
                or ("ULTRA" in line_upper and "ultra" in species_id)
                or ("WHITE" in line_upper and species_id == "kyuremwhite")
                or ("DUSK" in line_upper and species_id == "necrozmaduskmane")
                or ("WINGS" in line_upper and species_id == "necrozmadawnwings")
            ):
                choice = line
                break

    choice = re.sub(r"\s*\([^)]*\)", "", choice).strip()
    if " OR " in choice.upper():
        choice = choice.split(" OR ", 1)[0].strip()
    choice = choice.replace("!", "")
    return resolve_name(choice, canonical_maps["abilities"])


def build_set(block, species_name, canonical_maps):
    moves = [normalize_move(move, canonical_maps) for move in block["moves"]]
    set_data = {
        "level": block["level"],
        "ability": normalize_ability(block["ability"], species_name, canonical_maps),
        "moves": moves,
        "nature": block["nature"],
        "index": block["slot_index"],
        "teamIndex": block["team_index"],
    }
    if block.get("trainer_side"):
        set_data["trainerSide"] = block["trainer_side"]

    item = normalize_item(block["item"], canonical_maps)
    if item:
        set_data["item"] = item

    hidden_power_ivs = {}
    for move in moves:
        if move.startswith("Hidden Power "):
            hp_type = move.split("Hidden Power ", 1)[1].upper()
            hidden_power_ivs.update(HIDDEN_POWER_IVS.get(hp_type, {}))
    if hidden_power_ivs:
        set_data["ivs"] = hidden_power_ivs

    return set_data


def extract_blocks(workbook, order_caps):
    blocks = []
    encounter_index = 0
    for sheet_name in SHEETS:
        sheet = workbook[sheet_name]
        for row in range(1, sheet.max_row + 1):
            if not is_block_row(sheet, row):
                continue
            notes = extract_note_tags(sheet, row)
            base_label = build_base_label(sheet_name, sheet.cell(row, 3).value, notes)
            if not base_label:
                continue
            variant = note_variant(notes)
            members = []
            for col in SPECIES_COLS:
                species_value = sheet.cell(row, col).value
                if not isinstance(species_value, str) or not species_value.strip():
                    continue
                members.append(
                    {
                        "species": species_value.strip(),
                        "level": parse_level(sheet, sheet_name, base_label, row, col, order_caps),
                        "nature": " ".join(str(sheet.cell(row + 4, col).value).split()).strip(),
                        "ability": str(sheet.cell(row + 5, col).value).strip(),
                        "item": str(sheet.cell(row + 6, col).value).strip() if sheet.cell(row + 6, col).value else "",
                        "fillColor": cell_fill_color(sheet.cell(row, col)),
                        "moves": [
                            str(sheet.cell(move_row, col).value).strip()
                            for move_row in range(row + 7, row + 11)
                            if isinstance(sheet.cell(move_row, col).value, str) and str(sheet.cell(move_row, col).value).strip() != "-"
                        ],
                    }
                )
            if not members:
                continue
            assign_trainer_sides(members)
            blocks.append(
                {
                    "sheet": sheet_name,
                    "row": row,
                    "baseLabel": BLOCK_LABEL_ALIASES.get(base_label, base_label),
                    "variant": variant,
                    "encounterIndex": encounter_index,
                    "members": members,
                }
            )
            encounter_index += 1
    return blocks


def assign_variants(blocks):
    grouped = defaultdict(list)
    for block in blocks:
        grouped[block["baseLabel"]].append(block)

    for group in grouped.values():
        unnamed = [block for block in group if not block["variant"]]
        if len(group) == 1:
            continue
        if unnamed:
            for index, block in enumerate(sorted(unnamed, key=lambda value: value["encounterIndex"]), start=1):
                block["variant"] = f"TEAM {index}"

    for block in blocks:
        block["setName"] = block["baseLabel"] if not block["variant"] else f"{block['baseLabel']} [{block['variant']}]"

    names = [block["setName"] for block in blocks]
    duplicates = [name for name, count in Counter(names).items() if count > 1]
    if duplicates:
        raise ValueError(f"Duplicate set names remain: {duplicates}")


def sort_blocks(blocks, order_index):
    max_order = len(order_index)
    return sorted(
        blocks,
        key=lambda block: (
            order_index.get(block["baseLabel"], max_order + block["encounterIndex"]),
            block["encounterIndex"],
        ),
    )


def render_js_literal(value, indent=0):
    if isinstance(value, dict):
        if not value:
            return "{}"
        parts = []
        pad = " " * indent
        child_pad = " " * (indent + 2)
        for key, child in value.items():
            parts.append(f'{child_pad}"{key}": {render_js_literal(child, indent + 2)}')
        return "{\n" + ",\n".join(parts) + "\n" + pad + "}"
    if isinstance(value, list):
        return "[" + ", ".join(render_js_literal(item, indent) for item in value) + "]"
    return json.dumps(value, ensure_ascii=False)


def build_setdex(blocks, canonical_maps):
    setdex = {}
    for team_index, block in enumerate(blocks, start=1):
        for slot_index, member in enumerate(block["members"], start=1):
            species_name = normalize_species(member["species"], canonical_maps)
            set_data = build_set(
                {
                    "moves": member["moves"],
                    "ability": member["ability"],
                    "item": member["item"],
                    "nature": member["nature"],
                    "level": member["level"],
                    "trainer_side": member.get("trainerSide"),
                    "slot_index": slot_index,
                    "team_index": team_index,
                },
                species_name,
                canonical_maps,
            )
            setdex.setdefault(species_name, {})[block["setName"]] = set_data
    return dict(sorted(setdex.items(), key=lambda item: item[0]))


def validate_order(blocks, order_index):
    missing = sorted(
        {
            block["baseLabel"]
            for block in blocks
            if block["sheet"] != "Postgame" and block["baseLabel"] not in order_index
        }
    )
    if missing:
        raise ValueError("Missing main-game trainer order entries: " + ", ".join(missing))


def main():
    parser = argparse.ArgumentParser(description="Generate Hardcore trainer data from the Radical Red workbook.")
    parser.add_argument("workbook", help="Path to the Radical Red workbook")
    parser.add_argument(
        "--output",
        default="src/js/data/sets/hardcore.js",
        help="Relative output path for the generated hardcore set data",
    )
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parents[1]
    workbook = load_workbook(args.workbook, data_only=False)
    canonical_maps = load_canonical_maps(repo_root)
    order_entries = iter_order_entries(workbook)
    order_labels = [entry["label"] for entry in order_entries]
    order_index = {label: index for index, label in enumerate(order_labels)}
    order_caps = {entry["label"]: entry["cap"] for entry in order_entries if entry["cap"] is not None}
    blocks = extract_blocks(workbook, order_caps)
    assign_variants(blocks)
    validate_order(blocks, order_index)
    blocks = sort_blocks(blocks, order_index)
    setdex = build_setdex(blocks, canonical_maps)

    output_path = repo_root / args.output
    output_path.write_text(
        "var SETDEX_SV = " + render_js_literal(setdex, 0) + ";\n",
        encoding="utf-8",
    )

    print(
        json.dumps(
            {
                "teams": len(blocks),
                "trainerOrderEntries": len(order_labels),
                "species": len(setdex),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
